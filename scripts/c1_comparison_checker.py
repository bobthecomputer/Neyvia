"""Fresh disposable fixtures and independent read-only C1 comparison checks.

Only this module writes observation receipts. Call observe_fixture after each
revision; provider success text and action-return values are never accepted.
"""
from __future__ import annotations

import argparse
import gzip
import hashlib
import html
import json
from pathlib import Path
import secrets
import subprocess
import time
from xml.etree import ElementTree as ET

from run_c1_comparison import verify, DEFAULT, ROOT

PYTHON = 'C:/Users/user/AppData/Local/Programs/Python/Python313/python.exe'
EVIDENCE = ROOT / 'scripts/evidence'


def load(value):
    return json.loads(Path(value).read_text(encoding='utf-8')) if not isinstance(value, dict) else value


def safe_path(value):
    path = Path(value).resolve()
    if not any(path.is_relative_to(p) for p in (EVIDENCE, ROOT / '.agent_control')):
        raise ValueError('Checker refuses artifacts outside disposable task state')
    return path


def write_json(path, data):
    path.write_text(json.dumps(data, indent=2, ensure_ascii=False) + '\n', encoding='utf-8')


def prepare_fixtures(destination=EVIDENCE / 'C1CMP-fixtures'):
    verify(DEFAULT)
    frozen = load(DEFAULT)
    destination = safe_path(destination)
    claude = EVIDENCE / 'C1-claude-arm'
    claude.mkdir(exist_ok=True)
    index = []
    for task in frozen['tasks']:
        attempts = []
        for repetition in range(1, frozen['protocol']['repetitionsPerTask'] + 1):
            token = secrets.token_hex(16)
            for arm in frozen['protocol']['arms']:
                area = destination / arm / task['id'] / ('rep-' + str(repetition)) / token
                area.mkdir(parents=True, exist_ok=False)
                phrases = frozen['fixtures']['phrases']
                paths = {'root': str(area), 'profile': str(area / 'profile'),
                         'observations': str(area / 'observations.jsonl')}
                (area / 'profile').mkdir()
                (area / 'identity.txt').write_text(token + '\n' + '\n'.join(phrases) + '\n', encoding='utf-8')
                fixture_kind = task['fixture']
                if fixture_kind == 'editorFile':
                    paths['input'] = str(area / (token + '.txt'))
                    Path(paths['input']).write_text(token + '\n' + '\n'.join(phrases) + '\n', encoding='utf-8')
                elif fixture_kind == 'localForm':
                    paths['input'] = str(area / (token + '.html'))
                    Path(paths['input']).write_text('<!doctype html><meta charset="utf-8"><title>C1 ' + token + '</title>'
                        '<label for="note">C1 note</label><input id="note"><button onclick="document.querySelector(\'output\').textContent=document.querySelector(\'input\').value">Apply note</button><output></output>', encoding='utf-8')
                elif fixture_kind == 'renderedNote':
                    paths.update(input=str(area / 'note.html'), artifact=str(area / 'note.png'))
                    color = hashlib.sha256(token.encode()).hexdigest()[:6]
                    Path(paths['input']).write_text('<!doctype html><meta charset="utf-8"><title>C1 ' + token + '</title><style>html,body{margin:0;background:#' + color + ';color:white;font:32px Arial}</style><h1>' + token + '</h1>', encoding='utf-8')
                elif task['id'] == 'word-revise':
                    from docx import Document
                    paths['artifact'] = str(area / 'memo.docx')
                    document = Document(); document.add_paragraph(token); document.save(paths['artifact'])
                elif task['id'] == 'excel-edit':
                    from openpyxl import Workbook
                    paths['artifact'] = str(area / 'budget.xlsx')
                    book = Workbook(); sheet = book.active
                    sheet['A1'] = token; sheet['B1'] = 12; sheet['B2'] = 25; sheet['B3'] = '=SUM(B1:B2)'
                    book.save(paths['artifact'])
                elif task['id'] == 'powerpoint-slides':
                    from pptx import Presentation
                    paths['artifact'] = str(area / 'deck.pptx')
                    deck = Presentation()
                    for text in (token, 'Decisions', 'Next steps'):
                        slide = deck.slides.add_slide(deck.slide_layouts[5]); slide.shapes.title.text = text
                    deck.save(paths['artifact'])
                elif task['id'] == 'xournal-export':
                    paths.update(input=str(area / 'note.xopp'), artifact=str(area / 'note.pdf'))
                    note = ET.Element('xournal', creator='C1 disposable fixture', fileversion='4')
                    page = ET.SubElement(note, 'page', width='595', height='842')
                    ET.SubElement(page, 'background', type='solid', color='#ffffffff', style='plain')
                    layer = ET.SubElement(page, 'layer')
                    ET.SubElement(layer, 'text', font='Arial', size='16', x='36', y='48', color='#000000ff').text = token
                    with gzip.open(paths['input'], 'wb') as stream: stream.write(ET.tostring(note))
                elif fixture_kind == 'shellRoot':
                    paths.update(input=str(area / 'draft.txt'), artifact=str(area / 'persisted.txt'))
                    if task['id'] == 'git-note': paths['artifact'] = paths['input'] = str(area / 'note.txt')
                    Path(paths['input']).write_text(token + '\n', encoding='utf-8')
                    if task['id'] == 'git-note':
                        subprocess.run(['C:/Program Files/Git/cmd/git.exe', 'init', '--quiet', str(area)], check=True, creationflags=subprocess.CREATE_NO_WINDOW)
                record = {'id': task['id'], 'app': task['app'], 'arm': arm, 'repetition': repetition,
                    'token': token, 'instruction': task['instruction'], 'phrases': phrases,
                    'fixturePaths': paths, 'postcondition': task['postcondition'],
                    'timeBudgetSeconds': frozen['protocol']['timeBudgetSecondsPerAttempt'],
                    'freezeSha256': frozen['taskFreeze']['sha256']}
                manifest = area / 'fixture.json'
                record['manifest'] = str(manifest)
                record['checkerCommand'] = '& "' + PYTHON + '" "' + str(Path(__file__).resolve()) + '" --fixture "' + str(manifest) + '"'
                write_json(manifest, record)
                index.append(record)
                if arm == 'claude_computer_use': attempts.append(record)
        write_json(claude / (task['id'] + '.json'), {'id': task['id'], 'app': task['app'],
            'instruction': task['instruction'], 'timeBudgetSeconds': frozen['protocol']['timeBudgetSecondsPerAttempt'],
            'attempts': attempts, 'checkerUsage': 'Start checker with --watch before task; bind owned target HWND/control or CDP port to runtime in fixture.json for live-control tasks. Final checker never accepts claimed values.'})
    write_json(destination / 'index.json', {'freezeSha256': frozen['taskFreeze']['sha256'], 'attempts': index})
    return index


def read_artifact(fixture):
    f = load(fixture); paths = f['fixturePaths']; task = f['id']
    path = safe_path(paths.get('artifact', paths.get('input', paths['root'])))
    if task == 'word-revise':
        from docx import Document
        return '\n'.join(p.text for p in Document(path).paragraphs), {'path': str(path), 'source': 'fresh DOCX package read'}
    if task == 'excel-edit':
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=True); sheet = book.active
        return str(sheet['A1'].value), {'B3': sheet['B3'].value, 'path': str(path), 'source': 'fresh XLSX package read'}
    if task == 'powerpoint-slides':
        from pptx import Presentation
        deck = Presentation(path)
        return '\n'.join(s.text for s in deck.slides[0].shapes if s.has_text_frame), {'slideCount': len(deck.slides), 'path': str(path), 'source': 'fresh PPTX package read'}
    if task == 'xournal-export':
        from grant_agent import pdf_compat
        return '\n'.join(pdf_compat.page_texts(path)).strip(), {'pages': pdf_compat.page_count(path), 'path': str(path)}
    if task == 'git-note':
        def git(*args):
            return subprocess.check_output(['C:/Program Files/Git/cmd/git.exe', '-C', str(path.parent), *args], creationflags=subprocess.CREATE_NO_WINDOW, text=True).strip()
        return git('show', 'HEAD:' + path.name).splitlines()[0], {'cleanWorkingTree': git('status', '--porcelain') == '', 'remotes': git('remote'), 'commit': git('rev-parse', 'HEAD')}
    if task in ('cmd-note', 'powershell-note'):
        first = safe_path(paths['input']).read_text(encoding='utf-8-sig').strip('\r\n')
        second = path.read_text(encoding='utf-8-sig').strip('\r\n')
        return second, {'firstFile': first, 'files': 2}
    if task in ('firefox-render', 'zen-render'):
        from PIL import Image
        with Image.open(path) as image:
            rgb = image.convert('RGB'); color = tuple(bytes.fromhex(hashlib.sha256(f['phrases'][-1].encode()).hexdigest()[:6]))
            ink = sum(min(pixel) > 220 for pixel in rgb.crop((0, 0, min(rgb.width, 800), min(rgb.height, 220))).getdata())
            ok = rgb.width >= 800 and rgb.height >= 600 and rgb.getpixel((0, 0)) == color and ink > 1000
            return f['phrases'][-1] if ok else None, {'size': list(rgb.size), 'headingInkPixels': ink, 'cornerRgb': rgb.getpixel((0, 0)), 'boundary': 'Color and heading ink only; no OCR or producer attribution'}
    raise ValueError('Task needs live owned-control readback, not file inference')


def live_read(fixture, runtime=None):
    f = load(fixture); paths = f['fixturePaths']; task = f['id']
    if runtime and 'office' in runtime:
        doc = runtime['document']; app = runtime['office'].app_name
        if app == 'Microsoft Word': return str(doc.Content.Text).rstrip('\r'), {'source': 'independent Office COM Content.Text getter'}
        if app == 'Microsoft Excel': return str(doc.Worksheets(1).Range('A1').Value), {'source': 'independent Office COM Range getters', 'B3': float(doc.Worksheets(1).Range('B3').Value)}
        return str(doc.Slides(1).Shapes(1).TextFrame.TextRange.Text), {'source': 'independent Office COM TextRange getter', 'slideCount': doc.Slides.Count}
    if runtime and 'worker' in runtime:
        result = runtime['worker'].request('inspect', {'windowId': str(runtime['windowId'])})
        selector = runtime['selector']
        rows = [r for r in result['tree'] if all(str(r.get(k, '')) == str(v) for k, v in selector.items())]
        if len(rows) != 1: raise ValueError('Independent control target missing or ambiguous')
        return rows[0].get('value'), {'source': 'fresh owned-window inspect', 'windowId': str(runtime['windowId']), 'selector': selector, 'revision': result.get('revision')}
    binding = f.get('runtime', {})
    if binding.get('cdpPort'):
        from urllib.request import urlopen
        from urllib.parse import urlparse
        from websockets.sync.client import connect
        port = int(binding['cdpPort'])
        if not 48701 <= port <= 48709: raise ValueError('Checker CDP port outside task boundary')
        with urlopen(f'http://127.0.0.1:{port}/json/list', timeout=.5) as response: pages = json.load(response)
        uri = safe_path(paths['input']).as_uri()
        pages = [p for p in pages if p.get('type') == 'page' and (p.get('url') == uri or
            (p.get('url', '').startswith('vscode-file://vscode-app/') and f['token'] in p.get('title', '')))]
        if len(pages) != 1: raise ValueError('Independent CDP target absent or ambiguous')
        address = urlparse(pages[0]['webSocketDebuggerUrl'])
        if address.hostname != '127.0.0.1' or address.port != port: raise ValueError('CDP target identity mismatch')
        expression = "Array.from(document.querySelectorAll('input,textarea')).filter(e=>e.getBoundingClientRect().height>0).map(e=>({value:e.value,label:e.getAttribute('aria-label')||e.labels?.[0]?.textContent||e.id}))"
        with connect(pages[0]['webSocketDebuggerUrl'], open_timeout=.5, close_timeout=.2, proxy=None, origin=f'http://127.0.0.1:{port}') as socket:
            socket.send(json.dumps({'id': 1, 'method': 'Runtime.evaluate', 'params': {'expression': expression, 'returnByValue': True}}))
            while True:
                result = json.loads(socket.recv(timeout=.5))
                if result.get('id') == 1: break
        fields = result['result']['result']['value']
        label = binding.get('fieldLabel')
        if label: fields = [field for field in fields if field['label'] == label]
        if len(fields) != 1: raise ValueError('Independent CDP field absent or ambiguous; bind fieldLabel')
        return fields[0]['value'], {'source': 'fresh read-only CDP evaluation', 'port': port, 'url': pages[0]['url'], 'label': fields[0]['label']}
    if task == 'xournal-export':
        with gzip.open(safe_path(paths['input']), 'rb') as stream: value = ET.fromstring(stream.read()).find('page/layer/text').text
        return value, {'source': 'fresh gzip XML read', 'path': paths['input']}
    if task in ('firefox-render', 'zen-render'):
        import re
        text = safe_path(paths['input']).read_text(encoding='utf-8')
        match = re.search(r'<h1>(.*?)</h1>', text, re.S)
        return html.unescape(match.group(1)) if match else None, {'source': 'fresh HTML heading read', 'path': paths['input']}
    if task in ('cmd-note', 'powershell-note', 'git-note'):
        return safe_path(paths['input']).read_text(encoding='utf-8-sig').splitlines()[0], {'source': 'fresh disk read', 'path': paths['input']}
    if task.startswith(('word-', 'excel-', 'powerpoint-')): return read_artifact(f)
    raise ValueError('Live control binding absent; no fabricated sequence accepted')


def observe_fixture(fixture, phase=None, runtime=None):
    f = load(fixture)
    value, provenance = live_read(f, runtime)
    record = {'atUnixNs': time.time_ns(), 'phase': phase, 'value': value, 'provenance': provenance,
        'fixtureToken': f['token'], 'checkerSha256': hashlib.sha256(Path(__file__).read_bytes()).hexdigest()}
    with safe_path(f['fixturePaths']['observations']).open('a', encoding='utf-8') as stream:
        stream.write(json.dumps(record) + '\n')
    return record


def check_fixture(fixture, observations=None):
    f = load(fixture); expected = f['phrases']; condition = f['postcondition']; errors = []
    observed_path = safe_path(f['fixturePaths']['observations'])
    records = [json.loads(line) for line in observed_path.read_text(encoding='utf-8').splitlines()] if observed_path.exists() else []
    # Never accept caller-supplied observations; disk receipts contain checker getter provenance.
    if observations is not None: raise ValueError('Caller-supplied observations are not independent evidence')
    digest = hashlib.sha256(Path(__file__).read_bytes()).hexdigest()
    records = [r for r in records if r.get('fixtureToken') == f['token'] and r.get('checkerSha256') == digest and r.get('provenance', {}).get('source')]
    sequence = []
    for record in records:
        if not sequence or record['value'] != sequence[-1]: sequence.append(record['value'])
    sequence = [value for value in sequence if value in expected]
    requires_sequence = condition.get('freshReadbackEachRevision') or 'equalsOrdered' in condition
    if requires_sequence and sequence != expected: errors.append('Five ordered fresh independent revision readbacks missing')
    detail = {}
    try:
        if 'equalsOrdered' in condition:
            value = records[-1]['value'] if records else None
            if condition.get('sourceFileContainsAll'):
                text = safe_path(f['fixturePaths']['input']).read_text(encoding='utf-8')
                if not all(p in text for p in expected): errors.append('Editor fixture does not contain all phrases')
        else: value, detail = read_artifact(f)
        if value != expected[-1]: errors.append('Final value differs from frozen final phrase')
        for key, target in (('B3', 37), ('slideCount', 3), ('pages', 1), ('files', 2)):
            if key in detail and detail[key] != target: errors.append(key + ' differs from frozen postcondition')
        if 'firstFile' in detail and detail['firstFile'] != expected[-1]: errors.append('First file differs')
        if f['id'] == 'git-note' and (not detail.get('cleanWorkingTree') or detail.get('remotes')): errors.append('Git tree dirty or remote configured')
    except Exception as exc: errors.append(type(exc).__name__ + ': ' + str(exc))
    return {'ok': not errors, 'id': f['id'], 'arm': f['arm'], 'repetition': f['repetition'],
        'errors': errors, 'sequence': sequence, 'artifactReadback': detail, 'observationCount': len(records),
        'boundary': 'Artifact/control postconditions only; runner must separately verify app ownership, reopen, budget, guard and browser producer'}


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--prepare', action='store_true')
    parser.add_argument('--fixture', type=Path)
    parser.add_argument('--observe', action='store_true')
    parser.add_argument('--phase', type=int)
    parser.add_argument('--watch', action='store_true')
    args = parser.parse_args()
    if args.prepare: print(json.dumps({'prepared': len(prepare_fixtures())}))
    elif args.watch:
        f = load(args.fixture); deadline = time.monotonic() + f['timeBudgetSeconds']
        while time.monotonic() < deadline:
            try: observe_fixture(f)
            except Exception: pass
            time.sleep(.05)
        print(json.dumps(check_fixture(f)))
    elif args.observe: print(json.dumps(observe_fixture(args.fixture, args.phase)))
    else:
        result = check_fixture(args.fixture); print(json.dumps(result)); raise SystemExit(0 if result['ok'] else 1)
