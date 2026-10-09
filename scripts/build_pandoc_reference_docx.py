"""Build Neyvia's deterministic Pandoc DOCX reference document.

The document implements the compact_reference_guide design tokens used by the
managed Pandoc adapter.  Pandoc consumes its styles, numbering, and section
properties; the sample body is not copied into converted documents.
"""

from __future__ import annotations

import argparse
import os
import shutil
import subprocess
import tempfile
from datetime import datetime, timezone
from pathlib import Path
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


BLUE = RGBColor(0x2E, 0x74, 0xB5)
DARK_BLUE = RGBColor(0x1F, 0x4D, 0x78)
BLACK = RGBColor(0x00, 0x00, 0x00)
MUTED = RGBColor(0x66, 0x66, 0x66)
BUILD_TIME = datetime(2026, 7, 23, 0, 0, 0, tzinfo=timezone.utc)


def _set_font(style: object, name: str, size: float, color: RGBColor) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.color.rgb = color
    style.element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    style.element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)


def _configure_paragraph_style(
    style: object,
    *,
    font: str,
    size: float,
    color: RGBColor,
    before: float,
    after: float,
    line_spacing: float,
    bold: bool = False,
    keep_with_next: bool = False,
    keep_together: bool = False,
) -> None:
    _set_font(style, font, size, color)
    style.font.bold = bold
    paragraph = style.paragraph_format
    paragraph.space_before = Pt(before)
    paragraph.space_after = Pt(after)
    paragraph.line_spacing = line_spacing
    paragraph.keep_with_next = keep_with_next
    paragraph.keep_together = keep_together
    paragraph.widow_control = True


def _set_cell_margins(table_style: object) -> None:
    table_properties = table_style.element.find(qn("w:tblPr"))
    if table_properties is None:
        table_properties = OxmlElement("w:tblPr")
        table_style.element.append(table_properties)
    margins = table_properties.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        table_properties.append(margins)
    for side, value in (
        ("top", 80),
        ("start", 120),
        ("bottom", 80),
        ("end", 120),
    ):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def _set_table_borders(table_style: object) -> None:
    table_properties = table_style.element.find(qn("w:tblPr"))
    if table_properties is None:
        table_properties = OxmlElement("w:tblPr")
        table_style.element.append(table_properties)
    borders = table_properties.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        table_properties.append(borders)
    for edge in ("top", "start", "bottom", "end", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "6")
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "D9E2F3")


def _clear_sample_body(document: Document) -> None:
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def _resolve_pandoc(value: str | None) -> Path:
    requested = str(value or os.environ.get("NEYVIA_PANDOC") or "").strip()
    if requested:
        candidate = Path(requested).expanduser().resolve()
        if candidate.is_file():
            return candidate
        raise FileNotFoundError(f"Pandoc executable does not exist: {candidate}")
    discovered = shutil.which("pandoc")
    if discovered:
        return Path(discovered).resolve()
    raise FileNotFoundError(
        "Pandoc is required to build its reference document. "
        "Pass --pandoc or set NEYVIA_PANDOC."
    )


def _write_default_reference(pandoc: Path, destination: Path) -> None:
    with destination.open("wb") as output:
        completed = subprocess.run(
            [str(pandoc), "--print-default-data-file=reference.docx"],
            stdout=output,
            stderr=subprocess.PIPE,
            check=False,
        )
    if completed.returncode != 0 or destination.stat().st_size < 1_000:
        detail = completed.stderr.decode("utf-8", errors="replace")[:2_000]
        raise RuntimeError(
            f"Pandoc default reference extraction failed: {detail}"
        )


def _normalize_package(source: Path, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    with ZipFile(source, "r") as reader, ZipFile(
        destination,
        "w",
        compression=ZIP_DEFLATED,
        compresslevel=9,
    ) as writer:
        for name in sorted(reader.namelist()):
            original = reader.getinfo(name)
            info = ZipInfo(name, (2026, 7, 23, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            writer.writestr(info, reader.read(name))


def build(output: Path, *, pandoc: Path) -> None:
    with tempfile.TemporaryDirectory(prefix="neyvia-pandoc-base-") as temp_dir:
        default_reference = Path(temp_dir) / "pandoc-default-reference.docx"
        _write_default_reference(pandoc, default_reference)
        document = Document(default_reference)
    _clear_sample_body(document)
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = document.styles
    normal = styles["Normal"]
    _configure_paragraph_style(
        normal,
        font="Calibri",
        size=11,
        color=BLACK,
        before=0,
        after=6,
        line_spacing=1.25,
    )
    for style_name in ("Body Text", "First Paragraph", "Compact"):
        _configure_paragraph_style(
            styles[style_name],
            font="Calibri",
            size=11,
            color=BLACK,
            before=0,
            after=6,
            line_spacing=1.25,
            keep_together=style_name == "Compact",
        )
    _configure_paragraph_style(
        styles["Title"],
        font="Calibri",
        size=26,
        color=BLACK,
        before=0,
        after=8,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )
    _configure_paragraph_style(
        styles["Subtitle"],
        font="Calibri",
        size=13,
        color=MUTED,
        before=0,
        after=18,
        line_spacing=1.0,
    )
    _configure_paragraph_style(
        styles["Heading 1"],
        font="Calibri",
        size=16,
        color=BLUE,
        before=18,
        after=10,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )
    _configure_paragraph_style(
        styles["Heading 2"],
        font="Calibri",
        size=13,
        color=BLUE,
        before=14,
        after=7,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )
    _configure_paragraph_style(
        styles["Heading 3"],
        font="Calibri",
        size=12,
        color=DARK_BLUE,
        before=10,
        after=5,
        line_spacing=1.0,
        bold=True,
        keep_with_next=True,
    )
    for style_name in ("List Bullet", "List Number"):
        try:
            style = styles[style_name]
        except KeyError:
            continue
        _configure_paragraph_style(
            style,
            font="Calibri",
            size=11,
            color=BLACK,
            before=0,
            after=4,
            line_spacing=1.25,
            keep_together=True,
        )
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
    for style_name in ("Caption",):
        style = styles[style_name]
        _configure_paragraph_style(
            style,
            font="Calibri",
            size=9,
            color=MUTED,
            before=4,
            after=4,
            line_spacing=1.0,
        )
    if "Table Caption" not in styles:
        table_caption = styles.add_style(
            "Table Caption",
            WD_STYLE_TYPE.PARAGRAPH,
        )
        table_caption.base_style = styles["Caption"]
    table_style = styles["Table"]
    _set_font(table_style, "Calibri", 10, BLACK)
    _set_cell_margins(table_style)
    _set_table_borders(table_style)

    title = document.add_paragraph("Neyvia Reference Document", style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    document.add_paragraph(
        "Compact reference guide preset for deterministic Pandoc output.",
        style="Subtitle",
    )
    document.add_heading("Heading 1", level=1)
    document.add_paragraph("Body style and page geometry proof.")
    document.add_heading("Heading 2", level=2)
    document.add_paragraph("• List proof", style="Compact")
    document.add_heading("Heading 3", level=3)
    table = document.add_table(rows=2, cols=2, style="Table")
    table.autofit = False
    for row in table.rows:
        row.cells[0].width = Inches(1.875)
        row.cells[1].width = Inches(4.625)
    table.cell(0, 0).text = "Label"
    table.cell(0, 1).text = "Detail"
    table.cell(1, 0).text = "Geometry"
    table.cell(1, 1).text = "6.5 inch usable width"

    properties = document.core_properties
    properties.title = "Neyvia Pandoc Reference Document"
    properties.subject = "Deterministic document conversion styles"
    properties.author = "Neyvia"
    properties.last_modified_by = "Neyvia"
    properties.created = BUILD_TIME
    properties.modified = BUILD_TIME

    with tempfile.TemporaryDirectory(prefix="neyvia-pandoc-reference-") as temp_dir:
        raw = Path(temp_dir) / "reference.docx"
        document.save(raw)
        _normalize_package(raw, output)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("output")
    parser.add_argument("--pandoc")
    args = parser.parse_args()
    build(
        Path(args.output).resolve(),
        pandoc=_resolve_pandoc(args.pandoc),
    )
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
