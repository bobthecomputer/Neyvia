"""Fresh repair validation, kept separate from the frozen comparison denominator."""
import argparse
import copy
import json
from pathlib import Path
import uuid
from c1_comparison_execution import child_attempt, ROOT, save


def probe(task):
    index = json.loads((ROOT / 'scripts/evidence/C1CMP-fixtures/index.json').read_text(encoding='utf-8'))
    fixture = copy.deepcopy(next(f for f in index['attempts'] if f['id'] == task and f['arm'] == 'neyvia'))
    token = uuid.uuid4().hex
    app_folder = fixture['app'].replace(' ', '-') + '-' + token if task.endswith('-find') else token
    area = ROOT / '.agent_control/c1cmp-repair' / app_folder
    area.mkdir(parents=True)
    old = fixture['fixturePaths']
    paths = {'root': str(area), 'profile': str(area / 'profile'), 'observations': str(area / 'observations.jsonl')}
    (area / 'profile').mkdir()
    if old.get('input'):
        initial = Path(old['input']).read_bytes()
        # Independent initial document, same frozen phrases; new fixture ownership.
        paths['input'] = str(area / (token + Path(old['input']).suffix))
        Path(paths['input']).write_bytes(initial.replace(fixture['token'].encode(), token.encode()))
    if old.get('artifact'):
        paths['artifact'] = str(area / Path(old['artifact']).name)
        if Path(old['artifact']).exists():
            Path(paths['artifact']).write_bytes(Path(old['artifact']).read_bytes())
        if task == 'word-revise':
            from docx import Document
            document = Document(paths['artifact'])
            document.paragraphs[0].text = token
            document.save(paths['artifact'])
        elif task == 'excel-edit':
            from openpyxl import load_workbook
            book = load_workbook(paths['artifact'])
            book.active['A1'] = token
            book.save(paths['artifact'])
        elif task == 'powerpoint-slides':
            from pptx import Presentation
            deck = Presentation(paths['artifact'])
            deck.slides[0].shapes[0].text = token
            deck.save(paths['artifact'])
    if task == 'xournal-export':
        import gzip
        from xml.etree import ElementTree as ET
        with gzip.open(paths['input'], 'rb') as stream:
            note = ET.fromstring(stream.read())
        note.find('page/layer/text').text = token
        with gzip.open(paths['input'], 'wb') as stream:
            stream.write(ET.tostring(note))
    fixture.update(token=token, fixturePaths=paths, repetition=0, manifest=str(area / 'fixture.json'))
    path = area / 'fixture.json'
    save(path, fixture)
    destination = ROOT / 'scripts/evidence/C1CMP-repair' / (task + '-' + token + '.json')
    child_attempt(path, destination)
    result = json.loads(destination.read_text(encoding='utf-8'))
    result['comparisonScoreExcluded'] = True
    result['reason'] = 'Additional fresh repair validation; all frozen failed attempts remain recorded'
    save(destination, result)
    return result


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('task')
    args = parser.parse_args()
    probe(args.task)
